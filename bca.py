import re
from io import BytesIO

import numpy as np
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
import streamlit as st

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter


STANDARD_PATTERN = re.compile(
    r"^Std\s+([-+]?\d+(?:[.,]\d+)?)\s*ug/ml$",
    flags=re.IGNORECASE,
)

TECHNICAL_SUFFIX_PATTERN = re.compile(
    r"0001$",
    flags=re.IGNORECASE,
)


def clean_bca_sample_id(
    sample_id: str,
) -> str:
    """
    Usuwa techniczną końcówkę 0001 z nazw próbek.

    Przykłady:
    Std 1000 ug/ml0001 -> Std 1000 ug/ml
    Ctrl0001 -> Ctrl
    0,5 ug0001 -> 0,5 ug
    """

    cleaned_id = str(sample_id).strip()

    cleaned_id = TECHNICAL_SUFFIX_PATTERN.sub(
        "",
        cleaned_id,
    )

    return cleaned_id.strip()


def extract_standard_concentration(
    sample_id: str,
) -> float | None:
    """
    Wyciąga stężenie standardu z nazwy.

    Przykład:
    Std 62,5 ug/ml -> 62.5
    """

    match = STANDARD_PATTERN.match(
        str(sample_id).strip()
    )

    if match is None:
        return None

    concentration_text = (
        match.group(1)
        .replace(",", ".")
    )

    return float(concentration_text)


def extract_sample_dose(
    sample_id: str,
) -> float | None:
    """
    Wyciąga dawkę z nazwy próbki.

    Ctrl jest traktowany jako dawka 0.
    """

    normalized_id = (
        str(sample_id)
        .strip()
        .casefold()
    )

    if normalized_id in {
        "ctrl",
        "control",
        "kontrola",
    }:
        return 0.0

    concentration_text = (
        str(sample_id)
        .replace(",", ".")
    )

    match = re.search(
        r"([-+]?\d*\.?\d+)",
        concentration_text,
    )

    if match is None:
        return None

    return float(match.group(1))


def summarize_bca_absorbance(
    combined_data: pd.DataFrame,
) -> pd.DataFrame:
    """
    Czyści identyfikatory i oblicza:
    n, średnią, SD oraz CV absorbancji.
    """

    required_columns = {
        "sample_id",
        "raw_value",
    }

    missing_columns = (
        required_columns
        - set(combined_data.columns)
    )

    if missing_columns:
        raise ValueError(
            "Brakuje wymaganych kolumn: "
            + ", ".join(
                sorted(missing_columns)
            )
            + "."
        )

    data = combined_data.copy()

    data["sample_id"] = (
        data["sample_id"]
        .fillna("")
        .astype(str)
        .str.strip()
    )

    data = data.loc[
        data["sample_id"].ne("")
    ].copy()

    data["clean_sample_id"] = (
        data["sample_id"].map(
            clean_bca_sample_id
        )
    )

    data["raw_value"] = pd.to_numeric(
        data["raw_value"],
        errors="coerce",
    )

    data = data.dropna(
        subset=["raw_value"]
    )

    if data.empty:
        raise ValueError(
            "Brak opisanych wyników liczbowych "
            "do analizy BCA."
        )

    summary = (
        data
        .groupby(
            "clean_sample_id",
            as_index=False,
        )
        .agg(
            n=("raw_value", "count"),
            mean_absorbance=(
                "raw_value",
                "mean",
            ),
            sd_absorbance=(
                "raw_value",
                "std",
            ),
            minimum_absorbance=(
                "raw_value",
                "min",
            ),
            maximum_absorbance=(
                "raw_value",
                "max",
            ),
        )
        .rename(
            columns={
                "clean_sample_id": "sample_id",
            }
        )
    )

    summary["cv_pct"] = (
        summary["sd_absorbance"]
        / summary["mean_absorbance"].abs()
        * 100
    )

    summary.loc[
        summary["mean_absorbance"] == 0,
        "cv_pct",
    ] = pd.NA

    return summary


def classify_bca_results(
    absorbance_summary: pd.DataFrame,
) -> tuple[
    pd.DataFrame,
    pd.DataFrame,
    pd.DataFrame,
]:
    """
    Rozdziela wyniki na:
    - standardy,
    - blank,
    - właściwe próbki.
    """

    data = absorbance_summary.copy()

    data["standard_concentration"] = (
        data["sample_id"].map(
            extract_standard_concentration
        )
    )

    blank_mask = (
        data["sample_id"]
        .astype(str)
        .str.casefold()
        .str.startswith("blank")
    )

    standard_mask = (
        data["standard_concentration"]
        .notna()
    )

    blank_results = data.loc[
        blank_mask
    ].copy()

    standard_results = data.loc[
        standard_mask
    ].copy()

    sample_results = data.loc[
        ~blank_mask
        & ~standard_mask
    ].copy()

    if blank_results.empty:
        raise ValueError(
            "Nie znaleziono blanku. "
            "Oczekiwano identyfikatora "
            "rozpoczynającego się od 'Blank'."
        )

    if len(blank_results) > 1:
        raise ValueError(
            "Znaleziono więcej niż jeden "
            "identyfikator blanku."
        )

    if standard_results.empty:
        raise ValueError(
            "Nie znaleziono standardów BCA."
        )

    if standard_results[
        "standard_concentration"
    ].nunique() < 3:
        raise ValueError(
            "Do utworzenia krzywej potrzebne są "
            "co najmniej trzy różne standardy."
        )

    if sample_results.empty:
        raise ValueError(
            "Nie znaleziono właściwych próbek."
        )

    return (
        standard_results,
        blank_results,
        sample_results,
    )


def fit_bca_calibration(
    standard_results: pd.DataFrame,
    blank_results: pd.DataFrame,
) -> dict:
    """
    Odejmuje średni blank od standardów i dopasowuje:

    skorygowana absorbancja =
    slope * concentration + intercept
    """

    blank_mean = float(
        blank_results[
            "mean_absorbance"
        ].iloc[0]
    )

    standards = standard_results.copy()

    standards[
        "corrected_absorbance"
    ] = (
        standards["mean_absorbance"]
        - blank_mean
    )

    blank_standard = pd.DataFrame(
        [
            {
                "sample_id": "Blank",
                "n": int(
                    blank_results["n"].iloc[0]
                ),
                "mean_absorbance": blank_mean,
                "sd_absorbance": float(
                    blank_results[
                        "sd_absorbance"
                    ].iloc[0]
                ),
                "minimum_absorbance": float(
                    blank_results[
                        "minimum_absorbance"
                    ].iloc[0]
                ),
                "maximum_absorbance": float(
                    blank_results[
                        "maximum_absorbance"
                    ].iloc[0]
                ),
                "cv_pct": float(
                    blank_results[
                        "cv_pct"
                    ].iloc[0]
                ),
                "standard_concentration": 0.0,
                "corrected_absorbance": 0.0,
            }
        ]
    )

    calibration_data = pd.concat(
        [
            standards,
            blank_standard,
        ],
        ignore_index=True,
    )

    calibration_data = (
        calibration_data
        .sort_values(
            "standard_concentration"
        )
        .reset_index(drop=True)
    )

    x_values = calibration_data[
        "standard_concentration"
    ].to_numpy(
        dtype=float
    )

    y_values = calibration_data[
        "corrected_absorbance"
    ].to_numpy(
        dtype=float
    )

    slope, intercept = np.polyfit(
        x_values,
        y_values,
        1,
    )

    if slope <= 0:
        raise ValueError(
            "Nachylenie krzywej BCA "
            "musi być dodatnie."
        )

    predicted_values = (
        slope * x_values
        + intercept
    )

    residual_sum_of_squares = float(
        np.sum(
            (
                y_values
                - predicted_values
            ) ** 2
        )
    )

    total_sum_of_squares = float(
        np.sum(
            (
                y_values
                - np.mean(y_values)
            ) ** 2
        )
    )

    if total_sum_of_squares == 0:
        r_squared = np.nan
    else:
        r_squared = (
            1
            - residual_sum_of_squares
            / total_sum_of_squares
        )

    calibration_data[
        "calculated_concentration"
    ] = (
        calibration_data[
            "corrected_absorbance"
        ]
        - intercept
    ) / slope

    calibration_data[
        "difference_pct"
    ] = pd.NA

    nonzero_mask = (
        calibration_data[
            "standard_concentration"
        ] > 0
    )

    calibration_data.loc[
        nonzero_mask,
        "difference_pct",
    ] = (
        (
            calibration_data.loc[
                nonzero_mask,
                "calculated_concentration",
            ]
            - calibration_data.loc[
                nonzero_mask,
                "standard_concentration",
            ]
        )
        / calibration_data.loc[
            nonzero_mask,
            "standard_concentration",
        ]
        * 100
    )

    curve_x = np.linspace(
        0,
        float(
            calibration_data[
                "standard_concentration"
            ].max()
        ),
        300,
    )

    curve_y = (
        slope * curve_x
        + intercept
    )

    return {
        "calibration_data": calibration_data,
        "blank_mean": blank_mean,
        "slope": float(slope),
        "intercept": float(intercept),
        "r_squared": float(r_squared),
        "curve_x": curve_x,
        "curve_y": curve_y,
        "minimum_standard": float(
            calibration_data[
                "standard_concentration"
            ].min()
        ),
        "maximum_standard": float(
            calibration_data[
                "standard_concentration"
            ].max()
        ),
        "minimum_corrected_absorbance": float(
            calibration_data[
                "corrected_absorbance"
            ].min()
        ),
        "maximum_corrected_absorbance": float(
            calibration_data[
                "corrected_absorbance"
            ].max()
        ),
    }


def calculate_bca_samples(
    sample_results: pd.DataFrame,
    calibration_result: dict,
    dilution_factor: float,
    cv_limit: float,
) -> pd.DataFrame:
    """
    Oblicza końcowe stężenie białka dla próbek.
    """

    samples = sample_results.copy()

    blank_mean = calibration_result[
        "blank_mean"
    ]

    slope = calibration_result["slope"]
    intercept = calibration_result[
        "intercept"
    ]

    samples[
        "corrected_absorbance"
    ] = (
        samples["mean_absorbance"]
        - blank_mean
    )

    samples[
        "concentration_from_curve_ug_ml"
    ] = (
        samples[
            "corrected_absorbance"
        ]
        - intercept
    ) / slope

    samples["dilution_factor"] = (
        float(dilution_factor)
    )

    samples[
        "final_concentration_ug_ml"
    ] = (
        samples[
            "concentration_from_curve_ug_ml"
        ]
        * samples["dilution_factor"]
    )

    samples[
        "concentration_sd_ug_ml"
    ] = (
        samples["sd_absorbance"]
        / abs(slope)
        * samples["dilution_factor"]
    )

    samples["dose_ug"] = (
        samples["sample_id"].map(
            extract_sample_dose
        )
    )

    minimum_absorbance = (
        calibration_result[
            "minimum_corrected_absorbance"
        ]
    )

    maximum_absorbance = (
        calibration_result[
            "maximum_corrected_absorbance"
        ]
    )

    samples["range_status"] = "IN_RANGE"

    samples.loc[
        samples["corrected_absorbance"]
        < minimum_absorbance,
        "range_status",
    ] = "BELOW_RANGE"

    samples.loc[
        samples["corrected_absorbance"]
        > maximum_absorbance,
        "range_status",
    ] = "ABOVE_RANGE"

    samples["cv_status"] = np.where(
        samples["cv_pct"] <= cv_limit,
        "PASS",
        "HIGH_CV",
    )

    samples["overall_status"] = np.where(
        (
            samples["cv_status"].eq(
                "PASS"
            )
        )
        & (
            samples["range_status"].eq(
                "IN_RANGE"
            )
        ),
        "PASS",
        "CHECK",
    )

    return (
        samples
        .sort_values(
            [
                "dose_ug",
                "sample_id",
            ],
            na_position="last",
        )
        .reset_index(drop=True)
    )


def create_bca_calibration_figure(
    calibration_result: dict,
):
    """
    Tworzy liniową krzywą standardową BCA.
    """

    calibration_data = (
        calibration_result[
            "calibration_data"
        ]
    )

    figure = go.Figure()

    figure.add_trace(
        go.Scatter(
            x=calibration_data[
                "standard_concentration"
            ],
            y=calibration_data[
                "corrected_absorbance"
            ],
            error_y={
                "type": "data",
                "array": calibration_data[
                    "sd_absorbance"
                ],
                "visible": True,
            },
            mode="markers+text",
            text=calibration_data[
                "sample_id"
            ],
            textposition="top center",
            name="Standardy",
            marker={
                "size": 10,
                "color": "black",
            },
        )
    )

    figure.add_trace(
        go.Scatter(
            x=calibration_result[
                "curve_x"
            ],
            y=calibration_result[
                "curve_y"
            ],
            mode="lines",
            name="Regresja liniowa",
            line={
                "width": 3,
            },
        )
    )

    equation_text = (
        f"y = "
        f"{calibration_result['slope']:.6g}x "
        f"{calibration_result['intercept']:+.6g}"
        f"<br>R² = "
        f"{calibration_result['r_squared']:.4f}"
    )

    figure.add_annotation(
        x=0.06
        * calibration_result[
            "maximum_standard"
        ],
        y=0.90
        * float(
            np.nanmax(
                calibration_result[
                    "curve_y"
                ]
            )
        ),
        text=equation_text,
        showarrow=False,
        align="left",
        bgcolor="white",
        bordercolor="gray",
    )

    figure.update_layout(
        title="Krzywa standardowa BCA",
        xaxis_title=(
            "Stężenie białka [µg/mL]"
        ),
        yaxis_title=(
            "Absorbancja A562 po korekcji blanku"
        ),
        margin={
            "l": 70,
            "r": 40,
            "t": 80,
            "b": 70,
        },
    )

    return figure


def create_bca_sample_figure(
    sample_results: pd.DataFrame,
    chart_type: str,
):
    """
    Tworzy wykres końcowego stężenia białka.
    """

    chart_data = sample_results.copy()

    dose_order = (
        chart_data[
            [
                "sample_id",
                "dose_ug",
            ]
        ]
        .sort_values(
            "dose_ug",
            na_position="last",
        )
        ["sample_id"]
        .tolist()
    )

    if chart_type == "Punkty i linia":
        figure = px.line(
            chart_data,
            x="sample_id",
            y="final_concentration_ug_ml",
            error_y="concentration_sd_ug_ml",
            markers=True,
            category_orders={
                "sample_id": dose_order,
            },
            labels={
                "sample_id": "Próbka / dawka",
                "final_concentration_ug_ml": (
                    "Stężenie białka [µg/mL]"
                ),
            },
            title=(
                "Stężenie białka w próbkach"
            ),
        )

    else:
        figure = px.bar(
            chart_data,
            x="sample_id",
            y="final_concentration_ug_ml",
            error_y="concentration_sd_ug_ml",
            color="overall_status",
            category_orders={
                "sample_id": dose_order,
            },
            labels={
                "sample_id": "Próbka / dawka",
                "final_concentration_ug_ml": (
                    "Stężenie białka [µg/mL]"
                ),
                "overall_status": "Status",
            },
            title=(
                "Stężenie białka w próbkach"
            ),
        )

    figure.update_layout(
        xaxis_tickangle=-45,
    )

    return figure


def calculate_western_blot_preparation(
    sample_results: pd.DataFrame,
    protein_per_lane_ug: float,
    loading_volume_ul: float,
    total_mixture_volume_ul: float,
) -> pd.DataFrame:
    """
    Oblicza przygotowanie próbek do Western blot.

    Laemmli ma stężenie 2×, dlatego stanowi połowę
    końcowej objętości mieszaniny.
    """

    if protein_per_lane_ug <= 0:
        raise ValueError(
            "Ilość białka na studzienkę "
            "musi być większa od zera."
        )

    if loading_volume_ul <= 0:
        raise ValueError(
            "Objętość ładowana musi być "
            "większa od zera."
        )

    if total_mixture_volume_ul <= 0:
        raise ValueError(
            "Całkowita objętość mieszaniny "
            "musi być większa od zera."
        )

    if total_mixture_volume_ul < (
        loading_volume_ul
    ):
        raise ValueError(
            "Całkowita objętość mieszaniny "
            "nie może być mniejsza od objętości "
            "ładowanej na studzienkę."
        )

    preparation = sample_results[
        [
            "sample_id",
            "final_concentration_ug_ml",
        ]
    ].copy()

    preparation[
        "protein_per_lane_ug"
    ] = float(protein_per_lane_ug)

    preparation[
        "loading_volume_ul"
    ] = float(loading_volume_ul)

    preparation[
        "total_mixture_volume_ul"
    ] = float(total_mixture_volume_ul)

    preparation[
        "equivalent_lane_count"
    ] = (
        preparation[
            "total_mixture_volume_ul"
        ]
        / preparation[
            "loading_volume_ul"
        ]
    )

    preparation[
        "total_protein_required_ug"
    ] = (
        preparation[
            "protein_per_lane_ug"
        ]
        * preparation[
            "equivalent_lane_count"
        ]
    )

    preparation[
        "protein_concentration_ug_ul"
    ] = (
        preparation[
            "final_concentration_ug_ml"
        ]
        / 1000
    )

    preparation[
        "sample_volume_ul"
    ] = (
        preparation[
            "total_protein_required_ug"
        ]
        / preparation[
            "protein_concentration_ug_ul"
        ]
    )

    preparation[
        "laemmli_2x_volume_ul"
    ] = (
        preparation[
            "total_mixture_volume_ul"
        ]
        / 2
    )

    preparation["water_volume_ul"] = (
        preparation[
            "total_mixture_volume_ul"
        ]
        - preparation[
            "sample_volume_ul"
        ]
        - preparation[
            "laemmli_2x_volume_ul"
        ]
    )

    preparation["preparation_status"] = np.where(
        (
            preparation[
                "final_concentration_ug_ml"
            ] > 0
        )
        & (
            preparation[
                "water_volume_ul"
            ] >= 0
        ),
        "OK",
        "IMPOSSIBLE",
    )

    preparation[
        "status_message"
    ] = np.where(
        preparation[
            "preparation_status"
        ].eq("OK"),
        "Możliwe przygotowanie",
        (
            "Za niskie stężenie białka dla "
            "podanych parametrów"
        ),
    )

    return preparation


def style_excel_header(
    worksheet,
    row_number: int,
    column_count: int,
) -> None:
    """
    Formatuje wiersz nagłówka arkusza XLSX.
    """

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    bottom_border = Border(
        bottom=Side(
            style="thin",
            color="808080",
        )
    )

    for column_number in range(
        1,
        column_count + 1,
    ):
        cell = worksheet.cell(
            row=row_number,
            column=column_number,
        )

        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = bottom_border


def set_excel_column_widths(
    worksheet,
) -> None:
    """
    Dopasowuje szerokości kolumn.
    """

    for column_cells in worksheet.columns:
        maximum_length = 0

        column_letter = get_column_letter(
            column_cells[0].column
        )

        for cell in column_cells:
            value = cell.value

            if value is None:
                continue

            maximum_length = max(
                maximum_length,
                len(str(value)),
            )

        worksheet.column_dimensions[
            column_letter
        ].width = min(
            maximum_length + 3,
            34,
        )


def write_dataframe_to_sheet(
    worksheet,
    dataframe: pd.DataFrame,
    start_row: int = 1,
) -> None:
    """
    Zapisuje DataFrame do arkusza XLSX.
    """

    for column_number, column_name in enumerate(
        dataframe.columns,
        start=1,
    ):
        worksheet.cell(
            row=start_row,
            column=column_number,
            value=str(column_name),
        )

    style_excel_header(
        worksheet,
        start_row,
        len(dataframe.columns),
    )

    for row_number, row_values in enumerate(
        dataframe.itertuples(
            index=False,
            name=None,
        ),
        start=start_row + 1,
    ):
        for column_number, value in enumerate(
            row_values,
            start=1,
        ):
            if pd.isna(value):
                excel_value = None
            elif isinstance(
                value,
                np.generic,
            ):
                excel_value = value.item()
            else:
                excel_value = value

            worksheet.cell(
                row=row_number,
                column=column_number,
                value=excel_value,
            )


def create_bca_xlsx_report(
    calibration_result: dict,
    sample_results: pd.DataFrame,
    western_preparation: pd.DataFrame,
    dilution_factor: float,
    cv_limit: float,
) -> bytes:
    """
    Tworzy raport XLSX z analizą BCA
    i przygotowaniem próbek do Western blot.
    """

    workbook = Workbook()

    summary_sheet = workbook.active
    summary_sheet.title = "Summary"

    standards_sheet = workbook.create_sheet(
        "BCA Standards"
    )

    samples_sheet = workbook.create_sheet(
        "BCA Samples"
    )

    western_sheet = workbook.create_sheet(
        "Western Blot Preparation"
    )

    title_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7",
    )

    input_fill = PatternFill(
        fill_type="solid",
        fgColor="DDEBF7",
    )

    parameter_fill = PatternFill(
        fill_type="solid",
        fgColor="E4DFEC",
    )

    warning_fill = PatternFill(
        fill_type="solid",
        fgColor="FCE4D6",
    )

    summary_sheet["A1"] = "BCA analysis summary"
    summary_sheet["A1"].font = Font(
        bold=True,
        size=14,
    )
    summary_sheet["A1"].fill = title_fill

    summary_values = [
        (
            "Blank mean A562",
            calibration_result[
                "blank_mean"
            ],
        ),
        (
            "Regression slope",
            calibration_result["slope"],
        ),
        (
            "Regression intercept",
            calibration_result[
                "intercept"
            ],
        ),
        (
            "R squared",
            calibration_result[
                "r_squared"
            ],
        ),
        (
            "Minimum standard [ug/mL]",
            calibration_result[
                "minimum_standard"
            ],
        ),
        (
            "Maximum standard [ug/mL]",
            calibration_result[
                "maximum_standard"
            ],
        ),
        (
            "Sample dilution factor",
            dilution_factor,
        ),
        (
            "CV limit [%]",
            cv_limit,
        ),
        (
            "Number of samples",
            len(sample_results),
        ),
        (
            "Samples requiring review",
            int(
                sample_results[
                    "overall_status"
                ]
                .ne("PASS")
                .sum()
            ),
        ),
    ]

    for row_number, (
        parameter_name,
        parameter_value,
    ) in enumerate(
        summary_values,
        start=3,
    ):
        summary_sheet.cell(
            row=row_number,
            column=1,
            value=parameter_name,
        )

        summary_sheet.cell(
            row=row_number,
            column=2,
            value=parameter_value,
        )

        summary_sheet.cell(
            row=row_number,
            column=1,
        ).fill = parameter_fill

        summary_sheet.cell(
            row=row_number,
            column=2,
        ).fill = input_fill

    standards = calibration_result[
        "calibration_data"
    ].copy()

    write_dataframe_to_sheet(
        standards_sheet,
        standards,
    )

    write_dataframe_to_sheet(
        samples_sheet,
        sample_results,
    )

    write_dataframe_to_sheet(
        western_sheet,
        western_preparation,
    )

    for worksheet in (
        summary_sheet,
        standards_sheet,
        samples_sheet,
        western_sheet,
    ):
        worksheet.freeze_panes = "A2"
        worksheet.sheet_view.showGridLines = False
        set_excel_column_widths(
            worksheet
        )

    for row_number in range(
        2,
        western_sheet.max_row + 1,
    ):
        status_cell = western_sheet.cell(
            row=row_number,
            column=western_preparation.columns.get_loc(
                "preparation_status"
            ) + 1,
        )

        if status_cell.value == "IMPOSSIBLE":
            status_cell.fill = warning_fill

    output_stream = BytesIO()

    workbook.save(output_stream)

    output_stream.seek(0)

    return output_stream.getvalue()


def create_bca_conclusions(
    calibration_result: dict,
    sample_results: pd.DataFrame,
    cv_limit: float,
) -> list:
    """
    Tworzy krótkie podsumowanie analizy BCA.
    """

    conclusions = []

    conclusions.append(
        "Krzywa standardowa została opisana "
        f"równaniem y = "
        f"{calibration_result['slope']:.6g}x "
        f"{calibration_result['intercept']:+.6g}, "
        f"przy R² = "
        f"{calibration_result['r_squared']:.4f}."
    )

    conclusions.append(
        "Średnia absorbancja blanku wynosiła "
        f"{calibration_result['blank_mean']:.4f}."
    )

    high_cv_samples = (
        sample_results.loc[
            sample_results["cv_pct"]
            > cv_limit,
            "sample_id",
        ]
        .astype(str)
        .tolist()
    )

    if high_cv_samples:
        conclusions.append(
            f"Limit CV = {cv_limit:.1f}% "
            "przekroczyły próbki: "
            + ", ".join(high_cv_samples)
            + "."
        )
    else:
        conclusions.append(
            "Żadna próbka nie przekroczyła "
            f"limitu CV = {cv_limit:.1f}%."
        )

    out_of_range_samples = (
        sample_results.loc[
            ~sample_results[
                "range_status"
            ].eq("IN_RANGE"),
            "sample_id",
        ]
        .astype(str)
        .tolist()
    )

    if out_of_range_samples:
        conclusions.append(
            "Poza zakresem krzywej znalazły się: "
            + ", ".join(
                out_of_range_samples
            )
            + "."
        )
    else:
        conclusions.append(
            "Wszystkie próbki mieściły się "
            "w zakresie krzywej standardowej."
        )

    maximum_row = sample_results.loc[
        sample_results[
            "final_concentration_ug_ml"
        ].idxmax()
    ]

    conclusions.append(
        "Najwyższe końcowe stężenie białka "
        f"({float(maximum_row['final_concentration_ug_ml']):.2f} "
        "µg/mL) uzyskano dla próbki "
        f"{maximum_row['sample_id']}."
    )

    return conclusions


def show_bca_module(
    combined_data: pd.DataFrame,
) -> None:
    """
    Wyświetla kompletną analizę BCA
    i kalkulator Western blot.
    """

    st.subheader("Analiza BCA")

    st.write(
        "Standardy są korygowane o średni blank. "
        "Blank jest uwzględniany na krzywej jako "
        "punkt 0 µg/mL."
    )

    dilution_factor = st.number_input(
        "Współczynnik rozcieńczenia próbek",
        min_value=0.01,
        max_value=100000.0,
        value=20.0,
        step=1.0,
        key="bca_dilution_factor",
    )

    cv_limit = st.number_input(
        "Limit CV [%]",
        min_value=0.0,
        max_value=100.0,
        value=15.0,
        step=1.0,
        key="bca_cv_limit",
    )

    sample_chart_type = st.radio(
        "Rodzaj wykresu próbek",
        options=[
            "Punkty i linia",
            "Słupki",
        ],
        horizontal=True,
        key="bca_sample_chart_type",
    )

    try:
        absorbance_summary = (
            summarize_bca_absorbance(
                combined_data
            )
        )

        (
            standard_results,
            blank_results,
            raw_sample_results,
        ) = classify_bca_results(
            absorbance_summary
        )

        calibration_result = (
            fit_bca_calibration(
                standard_results,
                blank_results,
            )
        )

        sample_results = (
            calculate_bca_samples(
                raw_sample_results,
                calibration_result,
                dilution_factor,
                cv_limit,
            )
        )

        (
            metric_1,
            metric_2,
            metric_3,
            metric_4,
        ) = st.columns(4)

        metric_1.metric(
            "Średni blank A562",
            (
                f"{calibration_result['blank_mean']:.4f}"
            ),
        )

        metric_2.metric(
            "Nachylenie",
            (
                f"{calibration_result['slope']:.6g}"
            ),
        )

        metric_3.metric(
            "Wyraz wolny",
            (
                f"{calibration_result['intercept']:.6g}"
            ),
        )

        metric_4.metric(
            "R²",
            (
                f"{calibration_result['r_squared']:.4f}"
            ),
        )

        calibration_figure = (
            create_bca_calibration_figure(
                calibration_result
            )
        )

        st.plotly_chart(
            calibration_figure,
            width="stretch",
        )

        st.subheader(
            "Standardy i krzywa kalibracyjna"
        )

        st.dataframe(
            calibration_result[
                "calibration_data"
            ].style.format(
                {
                    "mean_absorbance": "{:.4f}",
                    "sd_absorbance": "{:.4f}",
                    "cv_pct": "{:.2f}",
                    "standard_concentration": (
                        "{:.4g}"
                    ),
                    "corrected_absorbance": (
                        "{:.4f}"
                    ),
                    "calculated_concentration": (
                        "{:.2f}"
                    ),
                    "difference_pct": "{:.2f}",
                },
                na_rep="brak",
            ),
            width="stretch",
            hide_index=True,
        )

        st.subheader(
            "Wyniki próbek"
        )

        st.dataframe(
            sample_results.style.format(
                {
                    "mean_absorbance": "{:.4f}",
                    "sd_absorbance": "{:.4f}",
                    "cv_pct": "{:.2f}",
                    "corrected_absorbance": (
                        "{:.4f}"
                    ),
                    "concentration_from_curve_ug_ml": (
                        "{:.2f}"
                    ),
                    "dilution_factor": "{:.2f}",
                    "final_concentration_ug_ml": (
                        "{:.2f}"
                    ),
                    "concentration_sd_ug_ml": (
                        "{:.2f}"
                    ),
                    "dose_ug": "{:.2f}",
                },
                na_rep="brak",
            ),
            width="stretch",
            hide_index=True,
        )

        sample_figure = (
            create_bca_sample_figure(
                sample_results,
                sample_chart_type,
            )
        )

        st.plotly_chart(
            sample_figure,
            width="stretch",
        )

        st.divider()

        st.subheader(
            "Przygotowanie próbek "
            "do Western blot"
        )

        st.info(
            "Obliczenia zakładają bufor "
            "Laemmli 2×, który stanowi połowę "
            "końcowej objętości mieszaniny."
        )

        preparation_column_1, (
            preparation_column_2
        ), preparation_column_3 = st.columns(
            3
        )

        protein_per_lane_ug = (
            preparation_column_1.number_input(
                "Białko na studzienkę [µg]",
                min_value=0.1,
                max_value=500.0,
                value=10.0,
                step=0.5,
                key="bca_wb_protein_per_lane",
            )
        )

        loading_volume_ul = (
            preparation_column_2.number_input(
                "Objętość ładowana [µL]",
                min_value=1.0,
                max_value=200.0,
                value=20.0,
                step=1.0,
                key="bca_wb_loading_volume",
            )
        )

        total_mixture_volume_ul = (
            preparation_column_3.number_input(
                "Całkowita objętość mieszaniny [µL]",
                min_value=1.0,
                max_value=5000.0,
                value=70.0,
                step=5.0,
                key="bca_wb_total_volume",
            )
        )

        western_preparation = (
            calculate_western_blot_preparation(
                sample_results,
                protein_per_lane_ug,
                loading_volume_ul,
                total_mixture_volume_ul,
            )
        )

        impossible_count = int(
            western_preparation[
                "preparation_status"
            ]
            .eq("IMPOSSIBLE")
            .sum()
        )

        if impossible_count > 0:
            st.warning(
                f"Dla {impossible_count} próbek "
                "nie można przygotować mieszaniny "
                "przy podanych parametrach."
            )

        st.dataframe(
            western_preparation.style.format(
                {
                    "final_concentration_ug_ml": (
                        "{:.2f}"
                    ),
                    "protein_per_lane_ug": (
                        "{:.2f}"
                    ),
                    "loading_volume_ul": (
                        "{:.2f}"
                    ),
                    "total_mixture_volume_ul": (
                        "{:.2f}"
                    ),
                    "equivalent_lane_count": (
                        "{:.2f}"
                    ),
                    "total_protein_required_ug": (
                        "{:.2f}"
                    ),
                    "protein_concentration_ug_ul": (
                        "{:.4f}"
                    ),
                    "sample_volume_ul": "{:.2f}",
                    "laemmli_2x_volume_ul": (
                        "{:.2f}"
                    ),
                    "water_volume_ul": "{:.2f}",
                },
                na_rep="brak",
            ),
            width="stretch",
            hide_index=True,
        )

        st.subheader(
            "Podsumowanie i krótkie wnioski"
        )

        conclusions = create_bca_conclusions(
            calibration_result,
            sample_results,
            cv_limit,
        )

        for conclusion in conclusions:
            st.markdown(
                f"- {conclusion}"
            )

        xlsx_report = create_bca_xlsx_report(
            calibration_result,
            sample_results,
            western_preparation,
            dilution_factor,
            cv_limit,
        )

        st.download_button(
            label=(
                "Pobierz analizę i przygotowanie "
                "próbek jako XLSX"
            ),
            data=xlsx_report,
            file_name=(
                "bca_western_blot_preparation.xlsx"
            ),
            mime=(
                "application/vnd.openxmlformats-"
                "officedocument.spreadsheetml.sheet"
            ),
        )

    except (
        ValueError,
        TypeError,
        ZeroDivisionError,
    ) as error:
        st.error(str(error))